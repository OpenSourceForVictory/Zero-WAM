"""Build auditable, provisional prompt and video mappings without changing datasets."""
import argparse,csv,json,hashlib
from pathlib import Path
from collections import Counter
import h5py,openpyxl,pyarrow.parquet as pq
from openpyxl.styles import Font,Alignment

def read(p):return json.loads(p.read_text())
def sheet(w,n):
 rows=list(w[n].values);return [dict(zip(rows[0],r)) for r in rows[1:] if any(v is not None for v in r)]
def write_csv(p,rows):
 with p.open('w',encoding='utf-8-sig',newline='') as f:
  writer=csv.DictWriter(f,fieldnames=list(rows[0]));writer.writeheader();writer.writerows(rows)
def main():
 ap=argparse.ArgumentParser();ap.add_argument('root',type=Path);ap.add_argument('output',type=Path);ap.add_argument('--order-confirmed',action='store_true');a=ap.parse_args();root=a.root;out=a.output;out.mkdir(parents=True,exist_ok=True)
 book=root/'Zero-WAM_ALOHA_长程任务_含12段验证视频_英文Prompt.xlsx';w=openpyxl.load_workbook(book,data_only=True)
 plan=sheet(w,'按布局采集96条');humanplan=sheet(w,'人类视频32段');valplan=sheet(w,'验证视频12段')
 datasets=list((root/'pika_aloha_10_20').glob('lerobot*/meta/source_episodes.json'));assert len(datasets)==1
 dataset=datasets[0].parent.parent;mapping=read(datasets[0])['episodes'];info=read(dataset/'meta/info.json');assert len(mapping)==len(plan)==96
 status='confirmed_recording_order' if a.order_confirmed else 'candidate_recording_order_needs_confirmation'
 def inventory(folder):
  rows=[]
  for d in sorted((root/folder/'human_video').iterdir()):
   if not d.is_dir():continue
   review=read(d/'review.json') if (d/'review.json').exists() else {};manifest=read(d/'manifest.json');v=d/'video.mp4'
   rows.append({'directory':d.name,'video':str(v.relative_to(root)),'accepted':review.get('accepted') is True,'reason':review.get('reason','missing_review'),'started_at':manifest['started_at'],'duration_s':manifest.get('finished_at',0)-manifest['started_at'],'exists':v.is_file(),'manifest_status':manifest.get('status')})
  return sorted(rows,key=lambda r:r['started_at'])
 inv=inventory('pika_aloha_10_20_human');valid=[r for r in inv if r['accepted']];assert len(valid)==len(humanplan)==32
 humans=[]
 for p,v in zip(humanplan,valid):
  assert v['exists'] and v['manifest_status']=='complete'
  humans.append({'human_demo_id':p['Human Demo ID'],'layout':p['布局ID'],'config':p['配置ID'],'sequence':p['顺序ID'],'human_video':v['video'],'prompt':p['English Filming Prompt'],'mapping_status':status,'content_review':'not_visually_verified'})
 lookup={r['human_demo_id']:r for r in humans};rows=[];current_tasks=Counter();frames=0
 for p,m in zip(plan,sorted(mapping,key=lambda r:r['episode_index'])):
  idx=m['episode_index'];assert p['采集序号']==idx+1
  hdf=root/'pika_aloha_10_20/hdf5'/m['source'];assert hdf.exists()
  with h5py.File(hdf) as f:
   assert len(f['action'])==m['length'];raw=str(f.attrs.get('source_episode',''));existing=str(f.attrs.get('task',''))
  table=pq.read_table(dataset/f'data/chunk-000/episode_{idx:06d}.parquet',columns=['episode_index','task_index']);assert table.num_rows==m['length'];assert set(table['episode_index'].to_pylist())=={idx};current_tasks.update(table['task_index'].to_pylist());frames+=m['length']
  human=lookup[p['人类视频ID']];assert human['prompt']==p['English Prompt'];assert human['layout']==p['布局ID']
  row={'episode_index':idx,'collection_number':p['采集序号'],'planned_episode_id':p['Episode ID'],'layout':p['布局ID'],'config':p['任务配置'],'sequence':p['顺序ID'],'repeat':p['重复'],'human_demo_id':p['人类视频ID'],'prompt':p['English Prompt'],'human_video':human['human_video'],'hdf5':str(hdf.relative_to(root)),'raw_source_recorded_path':raw,'frames':m['length'],'existing_hdf5_task':existing,'mapping_status':status,'content_review':'not_visually_verified'}
  for cam in ['cam_high','cam_left_wrist','cam_right_wrist']:
   v=dataset/f'videos/chunk-000/observation.images.{cam}/episode_{idx:06d}.mp4';assert v.exists();row[cam]=str(v.relative_to(root))
  rows.append(row)
 assert frames==info['total_frames'];assert set(Counter(r['human_demo_id'] for r in rows).values())=={3}
 vi=inventory('pika_aloha_10_20_human_verify');vv=[v for v in vi if v['accepted']];assert len(vv)==len(valplan)==12
 validations=[{'video_id':p['Video ID'],'dimension':p['Test dimension'],'pair':p['Matched pair'],'layout':p['Layout ID'],'prompt':p['Full English prompt'],'human_video':v['video'],'mapping_status':'candidate_recording_order_needs_separate_confirmation','robot_episode':'not_provided','split':'held_out_human_only'} for p,v in zip(valplan,vv)]
 excluded=[v for v in inv if not v['accepted']]
 for name,data in [('robot_human_pairs',rows),('human_demo_index',humans),('excluded_human_videos',excluded),('validation_human_index',validations)]:
  write_csv(out/(name+'.csv'),data);(out/(name+'.json')).write_text(json.dumps(data,ensure_ascii=False,indent=2))
 (out/'episode_prompts.candidate.json').write_text(json.dumps({str(r['episode_index']):r['prompt'] for r in rows},ensure_ascii=False,indent=2))
 summary={'source_workbook':str(book),'workbook_sha256':hashlib.sha256(book.read_bytes()).hexdigest(),'robot_episodes':len(rows),'frames':frames,'training_human_videos':len(humans),'excluded_human_videos':len(excluded),'validation_human_videos':len(validations),'distinct_prompts':len(set(r['prompt'] for r in rows)),'mapping_status':status,'original_lerobot_task_indices':dict(current_tasks),'original_tasks':(dataset/'meta/tasks.jsonl').read_text(),'datasets_modified':False,'content_visually_verified':False,'validation_layout_workbook_bytes':(root/'Zero-WAM_验证视频12段_含详细物体布局.xlsx').stat().st_size}
 (out/'audit.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2))
 export=openpyxl.Workbook();export.remove(export.active)
 for name,data in [('真机96条配对',rows),('有效人类32段',humans),('作废视频4段',excluded),('验证视频12段',validations)]:
  s=export.create_sheet(name);s.append(list(data[0]));
  for row in data:s.append(list(row.values()))
  s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
  for c in s[1]:c.font=Font(bold=True)
  for col in s.columns:s.column_dimensions[col[0].column_letter].width=50 if col[0].value in ['prompt','human_video'] else 23
  for row in s.iter_rows(min_row=2):
   for c in row:c.alignment=Alignment(vertical='top',wrap_text=True)
 export.save(out/'prompt与人类示教配对_待确认.xlsx');print(json.dumps(summary,ensure_ascii=False,indent=2))
if __name__=='__main__':main()
